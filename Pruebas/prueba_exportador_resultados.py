# ============================================================
# Archivo: prueba_exportador_resultados.py
# Proyecto: TecNet Finder
# Materia: Lenguajes y Autómatas II
#
# Descripción:
# Este archivo comprueba las exportaciones externas de los
# resultados actuales de TecNet Finder.
#
# Las pruebas verifican la creación del libro de Excel, la
# separación de redes en hojas y el contenido del reporte TXT.
# ============================================================

from openpyxl import load_workbook  # Importa la función utilizada para volver a abrir los libros creados durante las pruebas.
import pytest  # Importa pytest para comprobar errores esperados durante una exportación inválida.

from Motor.exportador_resultados import exportar_reporte_txt, exportar_vlsm_excel, generar_reporte_texto  # Importa las funciones públicas que serán verificadas.
from Motor.flujo_compilacion import FlujoCompilacion  # Importa el flujo real para producir resultados iguales a los utilizados por la interfaz.


def prueba_exporta_excel_con_una_red(tmp_path):  # Comprueba que una red válida produzca un libro con su tabla VLSM completa.
    flujo = FlujoCompilacion()  # Crea una instancia independiente del coordinador.
    codigo_fuente = "IP 192.168.1.0 MASK /24 HOSTS 50,30,10 NAME Oficina"  # Define una entrada válida que genera tres subredes.
    resultado = flujo.analizar(codigo_fuente)  # Ejecuta las fases actuales antes de exportar.
    ruta = tmp_path / "oficina.xlsx"  # Reserva una ubicación temporal para el archivo de prueba.

    exportar_vlsm_excel(resultado, ruta)  # Genera el libro utilizando el mismo módulo que empleará la interfaz.

    assert ruta.exists()  # Comprueba que el archivo haya sido creado físicamente.

    libro = load_workbook(ruta)  # Abre nuevamente el resultado para comprobar su contenido.
    assert libro.sheetnames == ["Oficina"]  # Verifica que el nombre de la red se utilice como nombre de hoja.

    hoja = libro["Oficina"]  # Obtiene la única hoja generada.
    assert hoja["A1"].value == "Red"  # Comprueba el primer encabezado de la tabla.
    assert hoja["B1"].value == "Hosts solicitados"  # Comprueba el encabezado correspondiente a la solicitud de hosts.
    assert hoja["A2"].value == "Oficina"  # Verifica que la primera fila conserve el nombre de la red.
    assert hoja["B2"].value == "50"  # Comprueba que la mayor solicitud aparezca primero según el cálculo VLSM.
    assert hoja.max_row == 4  # Verifica que existan una fila de encabezado y tres subredes.
    assert hoja.freeze_panes == "A2"  # Comprueba que el encabezado permanezca congelado al desplazarse.
    assert hoja.auto_filter.ref == "A1:I4"  # Comprueba que los filtros abarquen toda la tabla generada.


def prueba_exporta_excel_separa_redes_en_hojas(tmp_path):  # Comprueba que redes distintas se conserven en pestañas independientes.
    flujo = FlujoCompilacion()  # Crea el coordinador utilizado por la prueba.
    codigo_fuente = (  # Define dos instrucciones válidas con nombres diferentes.
        "IP 192.168.10.0 MASK /24 HOSTS 50 NAME Oficina\n"
        "IP 10.0.0.0 MASK /24 HOSTS 20 NAME Laboratorio"
    )
    resultado = flujo.analizar(codigo_fuente)  # Produce las subredes correspondientes a ambas instrucciones.
    ruta = tmp_path / "varias_redes.xlsx"  # Define la ubicación temporal del libro.

    exportar_vlsm_excel(resultado, ruta)  # Exporta las dos redes mediante el módulo real.

    libro = load_workbook(ruta)  # Abre el libro generado para inspeccionar sus hojas.
    assert libro.sheetnames == ["Oficina", "Laboratorio"]  # Comprueba que cada bloque tenga una pestaña propia.
    assert libro["Oficina"]["D2"].value == "192.168.10.0"  # Verifica que la primera hoja contenga la dirección de su red.
    assert libro["Laboratorio"]["D2"].value == "10.0.0.0"  # Verifica que la segunda hoja conserve su propia dirección.


def prueba_excel_rechaza_resultado_sin_subredes(tmp_path):  # Comprueba que no se produzca un libro vacío cuando la entrada no genera VLSM.
    flujo = FlujoCompilacion()  # Crea una instancia nueva del flujo.
    codigo_fuente = "IP 192.168.1. MASK /24 HOSTS 50"  # Define una IP incompleta que impide construir un bloque válido.
    resultado = flujo.analizar(codigo_fuente)  # Ejecuta el análisis para obtener un resultado sin subredes.
    ruta = tmp_path / "vacio.xlsx"  # Define una ubicación que no debería llegar a utilizarse.

    with pytest.raises(ValueError, match="No hay subredes VLSM"):  # Espera explícitamente el rechazo definido por el exportador.
        exportar_vlsm_excel(resultado, ruta)  # Intenta exportar un resultado que no contiene subredes.

    assert ruta.exists() is False  # Comprueba que el error no haya dejado un archivo vacío.


def prueba_reporte_texto_incluye_resultados_principales():  # Comprueba que el reporte reúna las secciones disponibles de la base actual.
    flujo = FlujoCompilacion()  # Crea una instancia independiente del coordinador.
    codigo_fuente = "IP 192.168.1.0 MASK /24 HOSTS 50 NAME Oficina"  # Define una instrucción válida con nombre.
    resultado = flujo.analizar(codigo_fuente)  # Produce tokens, bloque, árbol y tabla VLSM.

    reporte = generar_reporte_texto(codigo_fuente, resultado)  # Construye el contenido sin escribir todavía un archivo.

    assert "1. CÓDIGO FUENTE" in reporte  # Comprueba que se incluya la entrada original.
    assert codigo_fuente in reporte  # Verifica que el texto fuente aparezca exactamente dentro del reporte.
    assert "3. ANÁLISIS LÉXICO" in reporte  # Comprueba que se incluya la fase léxica.
    assert "4. ANÁLISIS SINTÁCTICO" in reporte  # Comprueba que se incluya la fase sintáctica.
    assert "6. PRODUCCIONES USADAS" in reporte  # Comprueba que se conserven las producciones generadas por el parser.
    assert "8. ÁRBOLES SINTÁCTICOS" in reporte  # Comprueba que se incluya la representación textual del árbol.
    assert "9. TABLA VLSM" in reporte  # Comprueba que el reporte finalice con los resultados de direccionamiento.
    assert "Oficina" in reporte  # Verifica que los datos concretos de la red permanezcan disponibles.


def prueba_reporte_texto_incluye_diagnosticos():  # Comprueba que una entrada con errores conserve los diagnósticos estructurados dentro del TXT.
    flujo = FlujoCompilacion()  # Crea una instancia nueva del flujo.
    codigo_fuente = "IP 192.168.1.0 MASK /24 HOSTS 50 NAME Oficina @"  # Agrega un símbolo que debe producir un error léxico.
    resultado = flujo.analizar(codigo_fuente)  # Ejecuta el flujo permitiendo que el diagnóstico quede registrado.

    reporte = generar_reporte_texto(codigo_fuente, resultado)  # Construye el reporte correspondiente a la entrada incorrecta.

    assert "LEX-002" in reporte  # Comprueba que aparezca el código estructurado del error léxico.
    assert "Solución:" in reporte  # Verifica que también se incluya la recomendación procedente del catálogo.


def prueba_exporta_reporte_txt_en_utf8(tmp_path):  # Comprueba que el reporte pueda guardarse y recuperarse conservando caracteres del español.
    flujo = FlujoCompilacion()  # Crea el coordinador utilizado por la prueba.
    codigo_fuente = "IP 10.0.0.0 MASK /24 HOSTS 20 NAME Laboratorio"  # Define una entrada válida sencilla.
    resultado = flujo.analizar(codigo_fuente)  # Produce el resultado que será exportado.
    ruta = tmp_path / "reporte.txt"  # Define una ubicación temporal para el archivo.

    exportar_reporte_txt(codigo_fuente, resultado, ruta)  # Escribe físicamente el reporte utilizando UTF-8.

    assert ruta.exists()  # Comprueba que el archivo haya sido creado.

    contenido = ruta.read_text(encoding="utf-8")  # Lee nuevamente el archivo con la misma codificación.
    assert "ANÁLISIS SINTÁCTICO" in contenido  # Verifica que los acentos se hayan conservado correctamente.
    assert "Laboratorio" in contenido  # Comprueba que los datos concretos de la entrada aparezcan en el archivo.