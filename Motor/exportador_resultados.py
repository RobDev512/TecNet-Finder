# ============================================================
# Archivo: exportador_resultados.py
# Proyecto: TecNet Finder
# Materia: Lenguajes y Autómatas II
#
# Descripción:
# Este archivo contiene funciones auxiliares para conservar los
# resultados actuales de TecNet Finder fuera de la interfaz.
#
# Permite exportar las subredes VLSM a un libro de Excel y
# generar un reporte TXT con la información producida por las
# fases léxica y sintáctica reconstruidas hasta el momento.
# ============================================================

from pathlib import Path  # Importa Path para manejar de forma uniforme las rutas de archivos recibidas.
import re  # Importa expresiones regulares para limpiar nombres de hojas de Excel.

from openpyxl import Workbook  # Importa la clase utilizada para crear libros de Excel.
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # Importa herramientas de formato visual para las celdas.
from openpyxl.utils import get_column_letter  # Importa una utilidad para convertir índices de columnas en letras.

from Motor.modelos import ResultadoCompilacion, SubredVLSM  # Importa los modelos necesarios para leer los resultados actuales del flujo.


ENCABEZADOS_VLSM = [  # Define el mismo orden de columnas utilizado actualmente por la interfaz.
    "Red",
    "Hosts solicitados",
    "Hosts disponibles",
    "Dirección de red",
    "CIDR",
    "Máscara",
    "Primera IP",
    "Última IP",
    "Broadcast",
]


def exportar_vlsm_excel(resultado: ResultadoCompilacion, ruta_archivo: str | Path) -> Path:  # Guarda las subredes VLSM agrupadas por bloque de red dentro de un libro de Excel.
    if not resultado.subredes_vlsm:  # Comprueba que exista al menos una subred calculada que pueda exportarse.
        raise ValueError("No hay subredes VLSM disponibles para exportar.")  # Evita crear un libro vacío cuando no existen resultados.

    ruta = Path(ruta_archivo)  # Convierte la ubicación recibida en un objeto Path.
    ruta.parent.mkdir(parents=True, exist_ok=True)  # Garantiza que la carpeta destino exista antes de guardar el archivo.

    libro = Workbook()  # Crea un libro nuevo utilizando openpyxl.
    hoja_inicial = libro.active  # Obtiene la hoja vacía creada automáticamente por openpyxl.
    libro.remove(hoja_inicial)  # Elimina la hoja predeterminada porque cada red tendrá su propia pestaña.

    grupos = _agrupar_subredes(resultado.subredes_vlsm)  # Separa las subredes según el bloque de red del que proceden.
    nombres_usados: set[str] = set()  # Conserva los nombres de hojas ya asignados para evitar duplicados.

    for (nombre_red, _ip_base), subredes in grupos.items():  # Recorre cada bloque manteniendo el orden de aparición del resultado.
        nombre_hoja = _crear_nombre_hoja(nombre_red, nombres_usados)  # Obtiene un nombre válido y único para la pestaña.
        hoja = libro.create_sheet(title=nombre_hoja)  # Crea una hoja independiente para el bloque actual.
        _llenar_hoja_vlsm(hoja, subredes)  # Inserta encabezados, datos y formato dentro de la nueva hoja.

    libro.save(ruta)  # Escribe físicamente el libro completo en la ubicación solicitada.
    return ruta  # Devuelve la ruta final para que la interfaz pueda informar dónde se guardó el archivo.


def generar_reporte_texto(codigo_fuente: str, resultado: ResultadoCompilacion) -> str:  # Construye el contenido completo del reporte TXT sin escribir todavía ningún archivo.
    lineas = [  # Inicia el reporte con un encabezado general fácilmente identificable.
        "TECNET FINDER - REPORTE DE ANÁLISIS",
        "=" * 38,
        "",
        "1. CÓDIGO FUENTE",
        "-" * 16,
        codigo_fuente if codigo_fuente else "(Entrada vacía)",
        "",
        "2. RESUMEN",
        "-" * 10,
        f"Tokens reconocidos: {len(resultado.resultado_lexico.tokens)}",
        f"Bloques sintácticamente válidos: {len(resultado.bloques_validos)}",
        f"Subredes VLSM generadas: {len(resultado.subredes_vlsm)}",
        f"Diagnósticos encontrados: {len(resultado.diagnosticos)}",
        "",
        "3. ANÁLISIS LÉXICO",
        "-" * 18,
    ]

    if resultado.resultado_lexico.tokens:  # Comprueba si el lexer produjo elementos que deban detallarse.
        for token in resultado.resultado_lexico.tokens:  # Recorre todos los tokens en el orden original.
            lineas.append(  # Agrega una línea compacta con tipo, valor y posición.
                f"{token.tipo:<12} | {token.valor:<18} | Línea {token.linea} | Columna {token.columna}"
            )
    else:  # Se ejecuta cuando no se reconoció ningún token.
        lineas.append("No se reconocieron tokens.")  # Deja constancia explícita de que la fase quedó vacía.

    lineas.extend(["", "4. ANÁLISIS SINTÁCTICO", "-" * 22])  # Abre la sección destinada a los bloques aceptados por el parser.

    if resultado.bloques_validos:  # Comprueba si existe al menos una instrucción sintácticamente válida.
        for indice, bloque in enumerate(resultado.bloques_validos, start=1):  # Recorre cada bloque válido conservando su orden.
            lineas.extend(  # Agrega los campos estructurados reconocidos por el parser.
                [
                    f"Bloque {indice}:",
                    f"  IP: {bloque.direccion_ip}",
                    f"  Máscara: {bloque.mascara_subred}",
                    f"  Hosts: {bloque.cantidades_hosts}",
                    f"  Nombre: {bloque.nombre or 'Sin nombre'}",
                    f"  Línea de origen: {bloque.linea_origen}",
                    "",
                ]
            )
    else:  # Se ejecuta cuando ninguna instrucción cumplió completamente la gramática.
        lineas.append("No se encontraron bloques sintácticamente válidos.")  # Informa el resultado sin inventar información inexistente.

    lineas.extend(["5. DIAGNÓSTICOS", "-" * 15])  # Abre la sección dedicada a errores y advertencias.

    if resultado.diagnosticos:  # Comprueba si alguna fase produjo diagnósticos.
        for indice, diagnostico in enumerate(resultado.diagnosticos, start=1):  # Recorre cada problema manteniendo su orden.
            ubicacion = _formatear_ubicacion_diagnostico(diagnostico.linea, diagnostico.columna)  # Convierte línea y columna en un texto legible.
            lineas.append(f"{indice}. [{diagnostico.codigo}] {diagnostico.categoria} - {diagnostico.severidad}")  # Identifica código, categoría y gravedad.
            lineas.append(f"   Ubicación: {ubicacion}")  # Agrega la posición relacionada con el problema.
            lineas.append(f"   Mensaje: {diagnostico.mensaje}")  # Conserva el mensaje específico generado durante el análisis.

            if diagnostico.fragmento:  # Comprueba si se almacenó el fragmento exacto relacionado con el diagnóstico.
                lineas.append(f"   Fragmento: {diagnostico.fragmento}")  # Incluye el texto problemático cuando está disponible.

            if diagnostico.solucion:  # Comprueba si el catálogo proporciona una recomendación.
                lineas.append(f"   Solución: {diagnostico.solucion}")  # Incluye la sugerencia de corrección correspondiente.

            lineas.append("")  # Separa visualmente diagnósticos consecutivos.
    else:  # Se ejecuta cuando la compilación actual no produjo problemas.
        lineas.extend(["No se encontraron diagnósticos.", ""])  # Registra explícitamente un resultado limpio.

    sintactico = resultado.resultado_sintactico  # Obtiene una referencia corta a la información adicional producida por el parser.

    lineas.extend(["6. PRODUCCIONES USADAS", "-" * 22])  # Abre la sección de gramática y producciones.
    lineas.append(sintactico.producciones_usadas or "No hay producciones disponibles.")  # Inserta el texto real generado por el analizador.

    lineas.extend(["", "7. DERIVACIÓN", "-" * 13])  # Abre la sección de derivación.
    lineas.append(sintactico.derivacion or "No hay derivación disponible.")  # Conserva la derivación producida por el parser.

    lineas.extend(["", "8. ÁRBOLES SINTÁCTICOS", "-" * 22])  # Abre la representación textual de los árboles.
    lineas.append(sintactico.arbol_textual or "No hay árboles sintácticos válidos para mostrar.")  # Agrega el árbol textual disponible.

    lineas.extend(["", "9. TABLA VLSM", "-" * 13])  # Abre la última sección correspondiente al cálculo actual.

    if resultado.subredes_vlsm:  # Comprueba si el flujo produjo resultados de direccionamiento.
        lineas.append(" | ".join(ENCABEZADOS_VLSM))  # Agrega los mismos encabezados utilizados por la interfaz y Excel.

        for subred in resultado.subredes_vlsm:  # Recorre todas las subredes calculadas.
            lineas.append(" | ".join(_valores_subred(subred)))  # Convierte cada resultado en una fila textual separada por barras.
    else:  # Se ejecuta cuando no fue posible calcular ninguna subred.
        lineas.append("No se generaron subredes VLSM.")  # Deja constancia explícita del resultado vacío.

    lineas.append("")  # Mantiene un salto final convencional en el archivo generado.
    return "\n".join(lineas)  # Une todas las secciones y devuelve el reporte completo listo para guardarse.


def exportar_reporte_txt(codigo_fuente: str, resultado: ResultadoCompilacion, ruta_archivo: str | Path) -> Path:  # Guarda en UTF-8 el reporte construido a partir de la compilación actual.
    ruta = Path(ruta_archivo)  # Convierte la ubicación recibida en un objeto Path.
    ruta.parent.mkdir(parents=True, exist_ok=True)  # Garantiza que la carpeta destino exista.
    contenido = generar_reporte_texto(codigo_fuente, resultado)  # Construye el texto utilizando únicamente datos estructurados del resultado.
    ruta.write_text(contenido, encoding="utf-8")  # Escribe el reporte conservando correctamente acentos y caracteres del español.
    return ruta  # Devuelve la ubicación final del archivo generado.


def _agrupar_subredes(subredes: list[SubredVLSM]) -> dict[tuple[str, str], list[SubredVLSM]]:  # Agrupa las subredes por nombre e IP base para separar correctamente bloques distintos.
    grupos: dict[tuple[str, str], list[SubredVLSM]] = {}  # Crea el diccionario que mantendrá el orden de aparición de cada bloque.

    for subred in subredes:  # Recorre los resultados VLSM en el orden producido por el flujo.
        clave = (subred.nombre_red, subred.ip_base)  # Usa nombre e IP base para evitar mezclar redes diferentes con un nombre repetido.
        grupos.setdefault(clave, []).append(subred)  # Crea el grupo cuando no existe y agrega después la subred actual.

    return grupos  # Devuelve todos los bloques organizados para construir las hojas del libro.


def _crear_nombre_hoja(nombre_red: str, nombres_usados: set[str]) -> str:  # Genera un nombre compatible con las restricciones de Excel.
    base = re.sub(r"[\[\]:*?/\\]", "_", nombre_red.strip()) or "Red"  # Sustituye caracteres prohibidos y evita nombres vacíos.
    candidato = base[:31]  # Limita inicialmente el nombre al máximo de 31 caracteres permitido por Excel.
    contador = 2  # Prepara el sufijo utilizado solamente cuando existe una hoja con el mismo nombre.

    while candidato in nombres_usados:  # Continúa hasta obtener una variante que todavía no se haya utilizado.
        sufijo = f"_{contador}"  # Construye un sufijo corto y predecible.
        candidato = f"{base[:31 - len(sufijo)]}{sufijo}"  # Reserva espacio para el sufijo sin superar el límite de Excel.
        contador += 1  # Incrementa el número para una posible siguiente colisión.

    nombres_usados.add(candidato)  # Registra el nombre finalmente seleccionado.
    return candidato  # Devuelve la etiqueta que puede utilizarse directamente como título de hoja.


def _llenar_hoja_vlsm(hoja, subredes: list[SubredVLSM]) -> None:  # Inserta y formatea una tabla VLSM completa dentro de una hoja.
    relleno_encabezado = PatternFill(fill_type="solid", fgColor="FF89B4FA")  # Utiliza el azul principal de TecNet Finder para los encabezados.
    fuente_encabezado = Font(bold=True, color="FF1E1E2E")  # Utiliza texto oscuro y negrita sobre el encabezado azul.
    borde_inferior = Border(bottom=Side(style="thin", color="FF45475A"))  # Añade una separación discreta entre filas.
    alineacion = Alignment(horizontal="center", vertical="center")  # Mantiene datos y encabezados centrados para facilitar su lectura.

    hoja.append(ENCABEZADOS_VLSM)  # Inserta la primera fila con las nueve columnas utilizadas por TecNet Finder.

    for celda in hoja[1]:  # Recorre exclusivamente las celdas de la fila de encabezado.
        celda.fill = relleno_encabezado  # Aplica el fondo azul principal.
        celda.font = fuente_encabezado  # Aplica negrita y color oscuro al texto.
        celda.alignment = alineacion  # Centra cada título dentro de su columna.

    for subred in subredes:  # Recorre todas las subredes pertenecientes al bloque actual.
        hoja.append(_valores_subred(subred))  # Inserta una fila utilizando el mismo orden de columnas definido en el encabezado.

    for fila in hoja.iter_rows(min_row=2):  # Recorre únicamente las filas que contienen resultados VLSM.
        for celda in fila:  # Recorre todas las columnas de la fila actual.
            celda.alignment = alineacion  # Centra el contenido para mantener una tabla uniforme.
            celda.border = borde_inferior  # Agrega una separación visual ligera entre resultados consecutivos.

    hoja.freeze_panes = "A2"  # Mantiene visible el encabezado al desplazarse por tablas extensas.
    hoja.auto_filter.ref = f"A1:I{hoja.max_row}"  # Habilita filtros sobre todas las columnas y filas generadas.
    hoja.sheet_view.showGridLines = False  # Oculta la cuadrícula predeterminada para que destaque el formato propio.
    hoja.row_dimensions[1].height = 24  # Da un poco más de altura al encabezado.

    _ajustar_ancho_columnas(hoja)  # Adapta cada columna al contenido real que fue insertado.


def _valores_subred(subred: SubredVLSM) -> list[str]:  # Convierte una subred en los nueve valores utilizados por las distintas exportaciones.
    return [  # Mantiene exactamente el mismo orden mostrado actualmente en la tabla de la interfaz.
        subred.nombre_red,
        str(subred.hosts_solicitados),
        str(subred.hosts_disponibles),
        subred.direccion_de_red,
        subred.nueva_mascara,
        subred.mascara_decimal,
        subred.primera_ip_utilizable,
        subred.ultima_ip_utilizable,
        subred.direccion_de_broadcast,
    ]


def _ajustar_ancho_columnas(hoja) -> None:  # Calcula un ancho legible para cada columna según su contenido.
    for indice_columna, columna in enumerate(hoja.iter_cols(), start=1):  # Recorre todas las columnas existentes en la hoja.
        ancho_contenido = max(len(str(celda.value or "")) for celda in columna)  # Obtiene la longitud máxima encontrada dentro de la columna.
        ancho_final = min(max(ancho_contenido + 3, 12), 32)  # Agrega espacio interior y evita columnas exageradamente estrechas o anchas.
        hoja.column_dimensions[get_column_letter(indice_columna)].width = ancho_final  # Aplica el ancho calculado utilizando la letra correspondiente.


def _formatear_ubicacion_diagnostico(linea: int | None, columna: int | None) -> str:  # Convierte coordenadas opcionales de un diagnóstico en una descripción legible.
    if linea is None:  # Comprueba si el diagnóstico no posee una línea concreta.
        return "Sin ubicación específica"  # Informa claramente que no existe una posición disponible.

    if columna is None:  # Comprueba si existe línea pero no columna.
        return f"Línea {linea}"  # Devuelve únicamente la línea conocida.

    return f"Línea {linea}, columna {columna}"  # Devuelve ambas coordenadas cuando están disponibles.